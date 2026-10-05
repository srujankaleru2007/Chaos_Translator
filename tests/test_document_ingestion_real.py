from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.documents import convert_document_bytes


def _converter_for(filename):
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions

    if not filename.endswith(".pdf"):
        return DocumentConverter()
    options = PdfPipelineOptions(do_ocr=False, do_table_structure=False, force_backend_text=True)
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET\n".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def _minimal_docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("DOCX provenance heading", level=1)
    document.add_paragraph("DOCX provenance text")
    document.save(output)
    return output.getvalue()


def _minimal_pptx() -> bytes:
    output = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX provenance heading"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "PPTX provenance text"
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "expected_text"),
    [("sample.pdf", "PDF provenance text"), ("sample.docx", "DOCX provenance text"), ("sample.pptx", "PPTX provenance text")],
)
def test_real_docling_conversion_preserves_content_and_provenance(filename, expected_text):
    docling = pytest.importorskip("docling")
    data = _minimal_pdf("PDF provenance text") if filename.endswith(".pdf") else _minimal_docx() if filename.endswith(".docx") else _minimal_pptx()
    source = convert_document_bytes(data, filename=filename, converter=_converter_for(filename))

    assert source.processing_status.value in {"completed", "partial"}
    assert source.extraction.tool == "docling"
    assert source.units
    if expected_text:
        assert any(expected_text in unit.text for unit in source.units)
    locator = source.units[0].locator
    if filename.endswith(".pptx"):
        assert locator.slide_number == 1 or locator.page_number == 1
    elif filename.endswith(".pdf"):
        assert locator.page_number == 1
    else:
        assert locator.metadata["docling_level"] >= 0
